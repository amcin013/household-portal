"""Reading a household spreadsheet, which is never shaped the way you'd hope."""
import io

import pytest
from openpyxl import Workbook, load_workbook

from portal.extensions import db
from portal.models import Account, Payee
from portal.services.importers import (
    export_workbook, guess_mapping, import_accounts, import_payees, parse_day,
    parse_money, preview_rows, read_table,
)


class Upload:
    """Stands in for a werkzeug FileStorage."""

    def __init__(self, data, filename):
        self._data = data
        self.filename = filename

    def read(self):
        return self._data


def sheet(rows, filename="book.xlsx"):
    wb = Workbook()
    ws = wb.active
    for row in rows:
        ws.append(row)
    buf = io.BytesIO()
    wb.save(buf)
    return Upload(buf.getvalue(), filename)


@pytest.mark.parametrize("raw,expected", [
    ("$145.50", 145.50),
    ("1,842.00", 1842.0),
    ("(650.00)", -650.0),          # accounting negative
    ("$1,234.56", 1234.56),
    (68.2, 68.2),
    ("", None),
    (None, None),
    ("n/a", None),
])
def test_parse_money(raw, expected):
    assert parse_money(raw) == expected


@pytest.mark.parametrize("raw,expected", [
    (1, 1), (15, 15), ("5th", 5), ("the 12th", 12),
    (45, None), (0, None), ("", None), (None, None),
])
def test_parse_day(raw, expected):
    assert parse_day(raw) == expected


def test_finds_header_row_past_a_title():
    upload = sheet([
        ["Monthly Household Budget 2026"],
        [],
        ["Company", "Amount", "Day"],
        ["TXU", 145.5, 1],
    ])
    headers, rows = read_table(upload)
    assert headers == ["Company", "Amount", "Day"]
    assert len(rows) == 1


def test_reads_a_plain_csv():
    data = b"Company,Amount,Day\nTXU,145.50,1\nWater,68.20,5\n"
    headers, rows = read_table(Upload(data, "bills.csv"))
    assert headers == ["Company", "Amount", "Day"]
    assert len(rows) == 2


def test_skips_blank_rows():
    upload = sheet([
        ["Company", "Amount"],
        ["TXU", 145.5],
        [None, None],
        ["", ""],
        ["Water", 68.2],
    ])
    _, rows = read_table(upload)
    assert len(rows) == 2


def test_guesses_columns_from_common_names():
    mapping = guess_mapping(["Company", "Amount Due", "Day of Month", "Type", "Website"])
    assert mapping["name"] == 0
    assert mapping["amount"] == 1
    assert mapping["due_day"] == 2
    assert mapping["category"] == 3
    assert mapping["url"] == 4


def test_guesses_alternative_column_names():
    mapping = guess_mapping(["Payee", "Monthly Payment", "Due Date"])
    assert mapping["name"] == 0
    assert mapping["amount"] == 1


def test_unrecognised_columns_are_simply_absent():
    mapping = guess_mapping(["Widget", "Sprocket"])
    assert "amount" not in mapping


def test_preview_skips_rows_with_no_name():
    headers = ["Company", "Amount"]
    rows = [["TXU", "145.50"], ["", "99.00"], ["Water", "68.20"]]
    items = preview_rows(headers, rows, {"name": 0, "amount": 1})
    assert [i["name"] for i in items] == ["TXU", "Water"]


def test_import_creates_payees(app):
    items = [
        {"name": "TXU", "amount": 145.5, "due_day": 1, "category": "Utilities",
         "url": "https://txu.com", "notes": "", "account_hint": "ends 4417"},
    ]
    created, updated = import_payees(items)
    assert (created, updated) == (1, 0)

    payee = Payee.query.filter_by(name="TXU").one()
    assert payee.typical_amount == 145.5
    assert payee.pay_url == "https://txu.com"
    assert payee.account_hint == "ends 4417"


def test_reimporting_updates_instead_of_duplicating(app):
    items = [{"name": "TXU", "amount": 145.5, "due_day": 1, "category": "Utilities",
              "url": "", "notes": "", "account_hint": ""}]
    import_payees(items)

    items[0]["amount"] = 162.0
    created, updated = import_payees(items)

    assert (created, updated) == (0, 1)
    assert Payee.query.filter_by(name="TXU").count() == 1
    assert Payee.query.filter_by(name="TXU").one().typical_amount == 162.0


def test_import_matches_names_case_insensitively(app):
    base = {"amount": 10, "due_day": 1, "category": "", "url": "", "notes": "",
            "account_hint": ""}
    import_payees([{**base, "name": "TXU Energy"}])
    import_payees([{**base, "name": "txu energy"}])
    assert Payee.query.count() == 1


def test_import_accounts_records_a_balance(app):
    items = [{"name": "TSP", "amount": 50000, "due_day": None, "category": "",
              "url": "", "notes": "", "account_hint": ""}]
    import_accounts(items, kind="retirement")

    acct = Account.query.filter_by(name="TSP").one()
    assert acct.kind == "retirement"
    assert acct.balance == 50000


def test_export_round_trips_through_openpyxl(app):
    db.session.add(Payee(name="TXU", category="Utilities", typical_amount=145.5,
                         due_day=1))
    db.session.commit()

    buf = export_workbook()
    wb = load_workbook(buf)

    assert {"Bills", "Accounts", "Recurring", "History"} <= set(wb.sheetnames)
    ws = wb["Bills"]
    assert ws.cell(row=1, column=1).value == "Company"
    assert ws.cell(row=2, column=1).value == "TXU"
    assert ws.cell(row=2, column=3).value == 145.5
